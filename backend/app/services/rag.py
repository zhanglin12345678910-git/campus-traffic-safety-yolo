from __future__ import annotations

import hashlib
import math
import re
import uuid
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from pypdf import PdfReader
from qdrant_client import QdrantClient, models

from app.config import Settings
from app.schemas import KnowledgeHit


class EmbeddingService:
    def __init__(self, settings: Settings, client: httpx.Client | None = None):
        self.settings = settings
        self._client = client

    def embed(self, texts: list[str]) -> list[list[float]]:
        if self.settings.embedding_provider == "openai_compatible":
            return self._remote_embed(texts)
        return [self._hash_embed(text) for text in texts]

    def _hash_embed(self, text: str) -> list[float]:
        """无需下载模型的中文字符 n-gram 向量，便于离线开发和可重复测试。"""
        dimension = self.settings.embedding_dimension
        vector = [0.0] * dimension
        normalized = re.sub(r"\s+", "", text.lower())
        tokens = [normalized[i : i + 2] for i in range(max(0, len(normalized) - 1))]
        tokens.extend(re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", text.lower()))
        for token in tokens:
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            index = int.from_bytes(digest[:4], "big") % dimension
            sign = 1.0 if digest[4] % 2 == 0 else -1.0
            vector[index] += sign
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]

    def _remote_embed(self, texts: list[str]) -> list[list[float]]:
        if not self.settings.embedding_base_url or not self.settings.embedding_api_key:
            raise RuntimeError("远程 Embedding 未完整配置")
        url = self.settings.embedding_base_url.rstrip("/")
        if not url.endswith("/embeddings"):
            url = f"{url}/embeddings"
        client = self._client or httpx.Client(
            timeout=self.settings.llm_timeout_seconds,
            proxy=self.settings.outbound_http_proxy,
        )
        close_client = self._client is None
        try:
            response = client.post(
                url,
                headers={"Authorization": f"Bearer {self.settings.embedding_api_key}"},
                json={"model": self.settings.embedding_model, "input": texts},
            )
            response.raise_for_status()
            data = sorted(response.json()["data"], key=lambda item: item["index"])
            return [list(item["embedding"]) for item in data]
        finally:
            if close_client:
                client.close()


class KnowledgeBaseService:
    def __init__(
        self,
        settings: Settings,
        client: QdrantClient | None = None,
        embedder: EmbeddingService | None = None,
    ):
        self.settings = settings
        self.embedder = embedder or EmbeddingService(settings)
        self._client = client
        self._collection_ready = False

    @property
    def client(self) -> QdrantClient:
        if self._client is None:
            if self.settings.qdrant_url:
                self._client = QdrantClient(
                    url=self.settings.qdrant_url,
                    api_key=self.settings.qdrant_api_key,
                    timeout=5,
                )
            else:
                self._client = QdrantClient(path=self.settings.qdrant_path)
        return self._client

    def health(self) -> dict[str, Any]:
        try:
            self.client.get_collections()
            return {"status": "ok", "mode": "server" if self.settings.qdrant_url else "embedded"}
        except Exception as exc:
            return {"status": "unavailable", "message": str(exc)}

    def ingest(self, document_id: str, name: str, path: str | Path) -> int:
        text = self.extract_text(path)
        chunks = self.chunk_text(text)
        if not chunks:
            raise ValueError("文档没有可入库的文本内容")
        vectors = self.embedder.embed(chunks)
        self._ensure_collection(len(vectors[0]))
        points = []
        for index, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
            chunk_id = f"{document_id}:{index}"
            point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, chunk_id))
            points.append(
                models.PointStruct(
                    id=point_id,
                    vector=vector,
                    payload={
                        "document_id": document_id,
                        "document_name": name,
                        "chunk_id": chunk_id,
                        "content": chunk,
                    },
                )
            )
        self.client.upsert(collection_name=self.settings.qdrant_collection, points=points, wait=True)
        return len(points)

    def delete_document(self, document_id: str) -> int:
        """按 document_id 删除该文档在向量库中的全部分块，返回删除前的分块数。

        用于替换写错的知识文档，以及清理历史测试反复上传的演示文档。
        嵌入模式（qdrant_path）下向量目录被当前进程独占，无法从外部脚本改写，
        因此删除必须经由后端进程执行。
        """
        self._ensure_collection(self.settings.embedding_dimension)
        selector = models.Filter(
            must=[
                models.FieldCondition(
                    key="document_id",
                    match=models.MatchValue(value=document_id),
                )
            ]
        )
        existing = self.client.count(
            collection_name=self.settings.qdrant_collection,
            count_filter=selector,
            exact=True,
        ).count
        if existing:
            self.client.delete(
                collection_name=self.settings.qdrant_collection,
                points_selector=selector,
                wait=True,
            )
        return int(existing)

    def search(self, query: str, top_k: int | None = None) -> list[KnowledgeHit]:
        if not query.strip():
            return []
        vector = self.embedder.embed([query])[0]
        self._ensure_collection(len(vector))
        result = self.client.query_points(
            collection_name=self.settings.qdrant_collection,
            query=vector,
            with_payload=True,
            limit=top_k or self.settings.rag_top_k,
            score_threshold=self.settings.rag_score_threshold,
        )
        hits = []
        for point in result.points:
            payload = point.payload or {}
            hits.append(
                KnowledgeHit(
                    document_name=str(payload.get("document_name", "未知文档")),
                    chunk_id=str(payload.get("chunk_id", point.id)),
                    content=str(payload.get("content", "")),
                    score=float(point.score),
                )
            )
        return hits

    def _ensure_collection(self, vector_size: int) -> None:
        if self._collection_ready:
            return
        if not self.client.collection_exists(self.settings.qdrant_collection):
            self.client.create_collection(
                collection_name=self.settings.qdrant_collection,
                vectors_config=models.VectorParams(size=vector_size, distance=models.Distance.COSINE),
            )
        self._collection_ready = True

    def extract_text(self, path: str | Path) -> str:
        source = Path(path)
        suffix = source.suffix.lower()
        if suffix in {".txt", ".md", ".markdown"}:
            return source.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".pdf":
            return "\n".join((page.extract_text() or "") for page in PdfReader(str(source)).pages)
        if suffix == ".docx":
            document = Document(str(source))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts)
        raise ValueError(f"不支持的知识文件类型：{suffix}")

    def chunk_text(self, text: str) -> list[str]:
        cleaned = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n")).strip()
        if not cleaned:
            return []
        size = self.settings.chunk_size
        overlap = min(self.settings.chunk_overlap, size - 1)
        chunks: list[str] = []
        start = 0
        while start < len(cleaned):
            end = min(len(cleaned), start + size)
            if end < len(cleaned):
                boundary = max(cleaned.rfind("\n", start, end), cleaned.rfind("。", start, end))
                if boundary > start + size // 2:
                    end = boundary + 1
            chunk = cleaned[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(cleaned):
                break
            start = max(start + 1, end - overlap)
        return chunks
